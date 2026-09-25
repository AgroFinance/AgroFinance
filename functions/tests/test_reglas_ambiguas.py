"""Regresión de 5 bugs reales encontrados con el dataset de prueba de una
agroexportadora real (mango + tuna): el motor clasificaba mal o ignoraba
consumos que sí traían todo el dato necesario, porque el tipo real vivía en
una columna vecina (o en una celda combinada de Word) que el motor no
miraba. Sin estos tests, un futuro cambio en ghg_classify/parse_archivo
puede reintroducir cualquiera de los 5 en silencio.
"""

from __future__ import annotations

import openpyxl

from engine.ghg_classify import ghg_classify
from engine.parse_archivo import parsear_archivo


def _escribir_xlsx(ruta, encabezado, filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(encabezado)
    for fila in filas:
        ws.append(fila)
    wb.save(str(ruta))


def test_combustible_por_fila_no_se_mezcla(tmp_path):
    """Bug real: 'Combustible_Litros' se cobraba SIEMPRE como diésel,
    ignorando que 'Tipo_Combustible' decía "gasohol" en algunas filas."""
    ruta = tmp_path / "combustible.xlsx"
    _escribir_xlsx(
        ruta,
        ["Equipo", "Combustible_Litros", "Tipo_Combustible"],
        [
            ["Tractor T-02", 22.5, "Diesel B5"],
            ["Camioneta V01", 8.6, "gasohol 84 (camioneta)"],
        ],
    )
    clasificadas = ghg_classify(parsear_archivo(str(ruta), "xlsx").lineas)
    factores = {round(l.valor, 1): l.factor_asignado for l in clasificadas}
    assert factores[22.5] == "dieselLitro"
    assert factores[8.6] == "gasohol84"


def test_gas_refrigerante_por_fila_no_se_mezcla(tmp_path):
    """Bug real: 'Gas_Recargado_kg' se ignoraba SIEMPRE ("gas no
    identificado"), aunque 'Gas_Refrigerante_Tipo' sí traía R-134a/R-404A
    por fila."""
    ruta = tmp_path / "refrigerante.xlsx"
    _escribir_xlsx(
        ruta,
        ["Equipo_ID", "Gas_Recargado_kg", "Gas_Refrigerante_Tipo"],
        [
            ["CF-01", 1.2, "R-134a"],
            ["CF-02", 2.0, "R-404A"],
        ],
    )
    clasificadas = ghg_classify(parsear_archivo(str(ruta), "xlsx").lineas)
    factores = {round(l.valor, 1): l.factor_asignado for l in clasificadas}
    assert factores[1.2] == "refrigeranteR134a"
    assert factores[2.0] == "refrigeranteR404a"


def test_columna_factor_emision_no_se_cuenta_como_consumo(tmp_path):
    """Bug real: 'Factor_Emision_Red_kgCO2_kWh' (el coeficiente, no un
    consumo) se clasificaba como electricidadSEIN solo por tener "kwh" en
    el nombre, inflando el total con una línea fantasma."""
    ruta = tmp_path / "energia.xlsx"
    _escribir_xlsx(
        ruta,
        ["kWh_Facturado", "Factor_Emision_Red_kgCO2_kWh"],
        [[17404.9, 0.2166]],
    )
    clasificadas = ghg_classify(parsear_archivo(str(ruta), "xlsx").lineas)
    assert len(clasificadas) == 2
    leida = next(l for l in clasificadas if l.estado == "leido")
    ignorada = next(l for l in clasificadas if l.estado == "ignorado")
    assert leida.factor_asignado == "electricidadSEIN"
    assert leida.valor == 17404.9
    assert ignorada.valor == 0.2166


def test_urea_aplicada_en_campo_es_alcance1_no_alcance3(tmp_path):
    """Bug real: una aplicación de urea EN CAMPO (con Metodo_Aplicacion/
    Superficie_ha/Parcela) se clasificaba como ureaProduccion (Alcance 3,
    compra) en vez de n2oSuelos (Alcance 1, emisión de suelo) — perdía su
    Alcance 1 por completo."""
    ruta = tmp_path / "aplicaciones.xlsx"
    _escribir_xlsx(
        ruta,
        ["Fundo", "Parcela_ID", "Tipo_Fertilizante", "Cantidad", "Unidad", "Metodo_Aplicacion", "Superficie_ha"],
        [["Fundo San Jose", "MG-01", "Urea", 60, "kg", "Fertirriego", 2.5]],
    )
    clasificadas = ghg_classify(parsear_archivo(str(ruta), "xlsx").lineas)
    urea = next(l for l in clasificadas if l.valor == 60)
    assert urea.factor_asignado == "n2oSuelos"
    assert urea.scope_asignado == 1


def test_urea_en_orden_de_compra_sigue_siendo_alcance3(tmp_path):
    """Contraparte del test anterior: una ORDEN DE COMPRA (sin columnas de
    aplicación en campo) debe seguir yendo a Alcance 3 — el fix de arriba
    no debe generalizarse a cualquier mención de urea."""
    ruta = tmp_path / "compras.xlsx"
    _escribir_xlsx(
        ruta,
        ["OC_Numero", "Proveedor", "Insumo", "Cantidad", "Unidad"],
        [["OC-001", "AgroQuimica SAC", "Urea 46%", 322, "kg"]],
    )
    clasificadas = ghg_classify(parsear_archivo(str(ruta), "xlsx").lineas)
    urea = next(l for l in clasificadas if l.valor == 322)
    assert urea.factor_asignado == "ureaProduccion"
    assert urea.scope_asignado == 3


def test_docx_celda_combinada_no_duplica_el_texto(tmp_path):
    """Bug real: una celda combinada (gridSpan) en una tabla Word hacía que
    python-docx devolviera el mismo texto dos veces por fila, y el
    extractor de texto libre contaba la recarga de refrigerante x2."""
    from docx import Document

    doc = Document()
    tabla = doc.add_table(rows=1, cols=2)
    tabla.cell(0, 0).merge(tabla.cell(0, 1))
    tabla.cell(0, 0).text = "Gas recargado R-404A, 1.8 kg"
    ruta = tmp_path / "orden_camara_fria.docx"
    doc.save(str(ruta))

    clasificadas = ghg_classify(parsear_archivo(str(ruta), "docx").lineas)
    leidas = [l for l in clasificadas if l.estado == "leido"]
    assert len(leidas) == 1
    assert leidas[0].factor_asignado == "refrigeranteR404a"
    assert leidas[0].valor == 1.8
